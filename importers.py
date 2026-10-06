"""Shared CSV/XLSX parsing logic for the Applications and Projects imports,
used by both the one-time CLI scripts (scripts/import_*.py) and the Admin
panel's web-based upload routes in app.py.
"""
import csv
from collections import namedtuple

from openpyxl import load_workbook

APPLICATIONS_FIELD_MAP = {
    "name": "name",
    "u_application_acronym": "acronym",
    "short_description": "description",
    "url": "url",
    "install_status": "status",
    "it_application_owner": "it_application_owner",
    "support_group": "support_group",
    "u_code_repository": "code_repository",
}
APPLICATIONS_REQUIRED_COLUMNS = ["name"]

PROJECTS_XLSX_FIELD_MAP = {
    "Short description": "short_description",
    "Federal Lead": "federal_lead",
    "Number": "number",
    "Configuration item": "configuration_item",
    "Stage": "stage",
    "State": "state",
    "Approval": "approval",
}
PROJECTS_XLSX_REQUIRED_COLUMNS = ["Number"]

# ServiceNow CSV export of sc_req_item with the OCIO Federal Lead catalog
# variable (identified by its sys_id) included as a column, covering all
# Federal Leads in one export instead of one pull per person.
PROJECTS_CSV_FIELD_MAP = {
    "number": "number",
    "short_description": "short_description",
    "variables.353504b61be56110f360a681f54bcbd5": "federal_lead",
    "stage": "stage",
    "state": "state",
    "approval": "approval",
    "cmdb_ci": "configuration_item",
}
PROJECTS_CSV_REQUIRED_COLUMNS = ["number"]

# records: list of field dicts ready for db.create_*(**fields)
# skipped: count of data rows dropped because a required column was blank
ImportResult = namedtuple("ImportResult", ["records", "skipped"])


class MissingColumnsError(ValueError):
    """Raised when an uploaded file is missing one or more required columns."""


def parse_applications_csv(file_obj):
    """file_obj: text-mode file-like object, e.g. a CMDB CSV export."""
    reader = csv.DictReader(file_obj)
    missing = [c for c in APPLICATIONS_REQUIRED_COLUMNS if c not in (reader.fieldnames or [])]
    if missing:
        raise MissingColumnsError(f"CSV is missing required column(s): {', '.join(missing)}")

    applications = []
    skipped = 0
    for row in reader:
        fields = {
            target: row[source].strip()
            for source, target in APPLICATIONS_FIELD_MAP.items()
            if row.get(source, "").strip()
        }
        if fields.get("name"):
            applications.append(fields)
        else:
            skipped += 1
    return ImportResult(applications, skipped)


def parse_projects_xlsx(file_obj):
    """file_obj: binary file-like object, e.g. an "ESRsWithLeads.xlsx" export."""
    workbook = load_workbook(file_obj, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]

    missing = [c for c in PROJECTS_XLSX_REQUIRED_COLUMNS if c not in header]
    if missing:
        raise MissingColumnsError(f"Spreadsheet is missing required column(s): {', '.join(missing)}")

    projects = []
    skipped = 0
    for row in rows:
        row_dict = dict(zip(header, row))
        fields = {
            target: str(row_dict[source]).strip()
            for source, target in PROJECTS_XLSX_FIELD_MAP.items()
            if row_dict.get(source) not in (None, "")
        }
        if fields.get("number"):
            projects.append(fields)
        else:
            skipped += 1
    return ImportResult(projects, skipped)


def parse_projects_csv(file_obj):
    """file_obj: text-mode file-like object, e.g. a ServiceNow sc_req_item
    CSV export (with the OCIO Federal Lead variable included by sys_id)."""
    reader = csv.DictReader(file_obj)
    missing = [c for c in PROJECTS_CSV_REQUIRED_COLUMNS if c not in (reader.fieldnames or [])]
    if missing:
        raise MissingColumnsError(f"CSV is missing required column(s): {', '.join(missing)}")

    projects = []
    skipped = 0
    for row in reader:
        fields = {
            target: row[source].strip()
            for source, target in PROJECTS_CSV_FIELD_MAP.items()
            if row.get(source, "").strip()
        }
        if fields.get("number"):
            projects.append(fields)
        else:
            skipped += 1
    return ImportResult(projects, skipped)
